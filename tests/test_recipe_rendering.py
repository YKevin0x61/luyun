from services.recipes.rendering import render_markdown_to_html
from services.recipes.sop_parse import ParsedRecipe
from services.recipes.structured_render import render_station_html, render_structured_recipe_body


def test_html_blocks_script_and_unsafe_style():
    html = render_markdown_to_html(
        "# T\n\n## S\n\n<script>alert(1)</script>"
        "<div style=\"background:url(javascript:alert(1))\">x</div>"
        "<a href=\"javascript:alert(1)\" target=\"_self\">bad</a>"
        "<a href=\"https://example.com\" target=\"_blank\">ok</a>"
    )
    assert "<script" not in html
    assert "javascript:" not in html
    assert "background" not in html
    assert 'href="https://example.com"' in html
    assert 'rel="noopener noreferrer"' in html


def test_html_stamps_recipe_id_on_wrapped_card():
    html = render_markdown_to_html(
        '# T\n\n## S\n\n<h3 class="recipe-title" data-recipe-id="42">艇仔粥</h3>\n\n比例 1:1\n'
    )
    assert 'data-recipe-id="42"' in html
    assert 'article' in html
    assert 'recipe-card' in html


def test_html_keeps_structured_ingredient_table_classes():
    html = render_markdown_to_html(
        '# T\n\n## S\n\n### 面团\n\n'
        '<table class="recipe-ingredients"><tbody>'
        '<tr><td class="recipe-ingredients-name">面粉</td>'
        '<td class="recipe-ingredients-amount">200 g</td></tr>'
        '</tbody></table>\n'
    )
    assert "recipe-ingredients" in html
    assert "recipe-ingredients-name" in html
    assert "recipe-ingredients-amount" in html
    assert "table-scroll" not in html


def test_html_strips_non_numeric_recipe_id():
    html = render_markdown_to_html(
        '# T\n\n## S\n\n<h3 class="recipe-title" data-recipe-id="nope">艇仔粥</h3>\n\n比例 1:1\n'
    )
    assert "data-recipe-id" not in html


def test_html_keeps_numeric_base_servings_qty():
    html = render_markdown_to_html(
        '# T\n\n## S\n\n<h3 class="recipe-title" data-recipe-id="1" '
        'data-base-servings-qty="4.5" data-base-servings-unit="人份">面团</h3>\n\nx\n'
    )
    assert 'data-base-servings-qty="4.5"' in html
    assert 'data-base-servings-unit="人份"' in html


def test_html_strips_non_numeric_base_servings_qty():
    html = render_markdown_to_html(
        '# T\n\n## S\n\n<h3 class="recipe-title" data-recipe-id="1" '
        'data-base-servings-qty="nope" data-base-servings-unit="人份">面团</h3>\n\nx\n'
    )
    assert "data-base-servings-qty" not in html
    assert "data-base-servings-unit" not in html


def test_html_wraps_table_in_scroll():
    html = render_markdown_to_html("## 配方\n\n| A | B |\n|:---|:---|\n| 1 | 2 |\n")
    assert "table-scroll" in html
    assert "<table" in html


def test_html_single_newline_renders_as_line_break():
    """正文内的单行换行应渲染为 <br>，与编辑器逐行罗列保持一致（nl2br）。"""
    html = render_markdown_to_html("## 配方\n\n### 馅\n\n盐：1克\n糖：2克\n")
    assert "<br" in html
    assert "盐：1克" in html
    assert "糖：2克" in html


def test_docx_renders_table_with_correct_dimensions():
    from services.recipes.rendering import render_markdown_to_docx

    md = "# 标题\n\n## 配方\n\n| 名称 | 用量 |\n|:---|:---|\n| 盐 | 1克 |\n| 糖 | 2克 |\n"
    doc = render_markdown_to_docx(md)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3
    assert len(table.columns) == 2
    assert table.rows[0].cells[0].text.strip() == "名称"
    assert table.rows[1].cells[0].text.strip() == "盐"


def test_docx_renders_page_break_div():
    from services.recipes.rendering import render_markdown_to_docx

    md = "# T\n\n## A\n\n正文一\n\n<div style=\"page-break-after: always;\"></div>\n\n## B\n\n正文二\n"
    doc = render_markdown_to_docx(md)
    xml = doc.element.xml
    assert 'w:type="page"' in xml


def test_station_docx_renders_ingredient_table_and_numbered_steps():
    from docx.shared import RGBColor

    from services.recipes.rendering import render_station_to_docx

    doc = render_station_to_docx("肠粉档", [
        ParsedRecipe(
            section="配方", recipe_name="面团", body_markdown="这段 Markdown 不应出现",
            sort_order=0, id=1,
            ingredients=({"name": "面粉", "amount": "200", "unit": "g"},),
            steps=("混合面粉与水", "静置 10 分钟"),
            tips=("夏天水温要更低",),
        ),
    ])
    assert len(doc.tables) == 1
    cell_texts = [cell.text.strip() for row in doc.tables[0].rows for cell in row.cells]
    assert "面粉" in cell_texts
    assert any("200" in text for text in cell_texts)
    paragraph_texts = [p.text for p in doc.paragraphs]
    assert not any("这段 Markdown 不应出现" in text for text in paragraph_texts)
    assert any(text.startswith("1.") and "混合面粉与水" in text for text in paragraph_texts)
    assert any(text.startswith("2.") and "静置 10 分钟" in text for text in paragraph_texts)
    tip_runs = [
        run for p in doc.paragraphs if "夏天水温要更低" in p.text for run in p.runs
    ]
    assert tip_runs
    assert any(run.font.color.rgb == RGBColor(180, 30, 30) for run in tip_runs)
    tip_paras = [p for p in doc.paragraphs if "夏天水温要更低" in p.text]
    assert tip_paras
    assert any('w:fill="FFF1F2"' in p._element.xml for p in tip_paras)


def test_station_docx_markdown_only_still_renders_body():
    from services.recipes.rendering import render_station_to_docx

    doc = render_station_to_docx("肠粉档", [
        ParsedRecipe(
            section="配方", recipe_name="肠粉酱油", body_markdown="酱油：100g",
            sort_order=0, id=1,
        ),
    ])
    assert len(doc.tables) == 0
    assert any("酱油：100g" in p.text for p in doc.paragraphs)


def test_structured_ingredients_empty_omits_table():
    assert render_structured_recipe_body([]) == ""
    assert render_structured_recipe_body(None) == ""


def test_structured_ingredients_one_row():
    html = render_structured_recipe_body(
        [{"name": "面粉", "amount": "200", "unit": "g"}],
    )
    assert html == (
        '<table class="recipe-ingredients"><tbody>'
        '<tr>'
        '<td class="recipe-ingredients-name">面粉</td>'
        '<td class="recipe-ingredients-amount">200 g</td>'
        '</tr>'
        '</tbody></table>'
    )


def test_structured_ingredients_many_rows():
    html = render_structured_recipe_body([
        {"name": "面粉", "amount": "200", "unit": "g"},
        {"name": "水", "amount": "120", "unit": "ml"},
        {"name": "盐", "amount": "2", "unit": "g"},
    ])
    assert html == (
        '<table class="recipe-ingredients"><tbody>'
        '<tr>'
        '<td class="recipe-ingredients-name">面粉</td>'
        '<td class="recipe-ingredients-amount">200 g</td>'
        '</tr>'
        '<tr>'
        '<td class="recipe-ingredients-name">水</td>'
        '<td class="recipe-ingredients-amount">120 ml</td>'
        '</tr>'
        '<tr>'
        '<td class="recipe-ingredients-name">盐</td>'
        '<td class="recipe-ingredients-amount">2 g</td>'
        '</tr>'
        '</tbody></table>'
    )


def test_structured_ingredients_non_numeric_amount():
    html = render_structured_recipe_body(
        [{"name": "胡椒", "amount": "适量", "unit": ""}],
    )
    assert html == (
        '<table class="recipe-ingredients"><tbody>'
        '<tr>'
        '<td class="recipe-ingredients-name">胡椒</td>'
        '<td class="recipe-ingredients-amount">适量</td>'
        '</tr>'
        '</tbody></table>'
    )


def test_structured_ingredients_escapes_html():
    html = render_structured_recipe_body(
        [{"name": "<script>x</script>", "amount": "1<", "unit": 'g>"'}],
    )
    assert "<script>" not in html
    assert "&lt;script&gt;" in html
    assert "1&lt;" in html
    assert "&quot;" in html


def test_station_html_keeps_markdown_and_structured_cards_in_one_grid():
    html = render_station_html("肠粉档", [
        ParsedRecipe(
            section="配方", recipe_name="肠粉酱油", body_markdown="酱油：100g",
            sort_order=0, id=1,
        ),
        ParsedRecipe(
            section="配方", recipe_name="面团", body_markdown="",
            sort_order=1, id=2,
            ingredients=({"name": "面粉", "amount": "200", "unit": "g"},),
        ),
    ])
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    grids = soup.select(".sop-section-grid")
    assert len(grids) == 1
    cards = grids[0].select("article.recipe-card")
    assert len(cards) == 2
    assert cards[0].get("data-recipe-id") == "1"
    assert cards[1].get("data-recipe-id") == "2"
    assert cards[0].select_one(".recipe-ingredients") is None
    assert "酱油：100g" in cards[0].select_one(".recipe-card-body").get_text()
    table = cards[1].select_one("table.recipe-ingredients")
    assert table is not None
    assert table.select_one(".recipe-ingredients-name").get_text() == "面粉"
    assert table.select_one(".recipe-ingredients-amount").get_text() == "200 g"


def test_station_html_markdown_only_still_builds_cards():
    html = render_station_html("肠粉档", [
        ParsedRecipe(
            section="配方", recipe_name="肠粉酱油", body_markdown="酱油：100g",
            sort_order=0, id=7,
        ),
    ])
    assert "recipe-card" in html
    assert 'data-recipe-id="7"' in html
    assert "recipe-ingredients" not in html
    assert "酱油：100g" in html


def test_station_html_escapes_recipe_name():
    html = render_station_html("肠粉档", [
        ParsedRecipe(
            section="配方", recipe_name="<img src=x>", body_markdown="x",
            sort_order=0, id=8,
        ),
    ])
    assert "<img" not in html
    assert "&lt;img src=x&gt;" in html


def test_structured_steps_empty_omits_list():
    assert render_structured_recipe_body(steps=[]) == ""
    assert render_structured_recipe_body(steps=None) == ""
    assert render_structured_recipe_body(ingredients=[], steps=[]) == ""


def test_structured_steps_one():
    html = render_structured_recipe_body(steps=["混合面粉与水"])
    assert html == (
        '<ol class="recipe-steps">'
        '<li class="recipe-steps-item">混合面粉与水</li>'
        '</ol>'
    )


def test_structured_steps_many():
    html = render_structured_recipe_body(steps=["混合面粉与水", "静置 10 分钟", "分成剂子"])
    assert html == (
        '<ol class="recipe-steps">'
        '<li class="recipe-steps-item">混合面粉与水</li>'
        '<li class="recipe-steps-item">静置 10 分钟</li>'
        '<li class="recipe-steps-item">分成剂子</li>'
        '</ol>'
    )


def test_structured_steps_escapes_html():
    html = render_structured_recipe_body(steps=['<script>x</script> 1<"'])
    assert "<script>" not in html
    assert "&lt;script&gt;" in html
    assert "&quot;" in html


def test_structured_body_concatenates_ingredients_then_steps():
    html = render_structured_recipe_body(
        ingredients=[{"name": "面粉", "amount": "200", "unit": "g"}],
        steps=["混合面粉与水"],
    )
    assert html == (
        '<table class="recipe-ingredients"><tbody>'
        '<tr>'
        '<td class="recipe-ingredients-name">面粉</td>'
        '<td class="recipe-ingredients-amount">200 g</td>'
        '</tr>'
        '</tbody></table>'
        '<ol class="recipe-steps">'
        '<li class="recipe-steps-item">混合面粉与水</li>'
        '</ol>'
    )


def test_html_keeps_structured_step_list_classes():
    html = render_markdown_to_html(
        '# T\n\n## S\n\n### 面团\n\n'
        '<ol class="recipe-steps">'
        '<li class="recipe-steps-item">混合面粉与水</li>'
        '</ol>\n'
    )
    assert "recipe-steps" in html
    assert "recipe-steps-item" in html


def test_station_html_steps_only_uses_structured_seam():
    html = render_station_html("肠粉档", [
        ParsedRecipe(
            section="配方", recipe_name="面团",
            body_markdown="这段 Markdown 不应出现在结构化卡片里",
            sort_order=0, id=9,
            steps=("混合面粉与水", "静置 10 分钟"),
        ),
    ])
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    card = soup.select_one("article.recipe-card")
    assert card is not None
    ol = card.select_one("ol.recipe-steps")
    assert ol is not None
    items = [li.get_text() for li in ol.select("li.recipe-steps-item")]
    assert items == ["混合面粉与水", "静置 10 分钟"]
    assert "这段 Markdown 不应出现在结构化卡片里" not in html
    assert card.select_one(".recipe-ingredients") is None


def test_station_html_keeps_ingredients_when_steps_present():
    html = render_station_html("肠粉档", [
        ParsedRecipe(
            section="配方", recipe_name="面团", body_markdown="",
            sort_order=0, id=10,
            ingredients=({"name": "面粉", "amount": "200", "unit": "g"},),
            steps=("混合面粉与水",),
        ),
    ])
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    body = soup.select_one(".recipe-card-body")
    assert body.select_one("table.recipe-ingredients") is not None
    assert body.select_one("ol.recipe-steps") is not None
    assert body.select_one(".recipe-ingredients-name").get_text() == "面粉"
    assert body.select_one(".recipe-steps-item").get_text() == "混合面粉与水"


def test_structured_tips_empty_omits_list():
    assert render_structured_recipe_body(tips=[]) == ""
    assert render_structured_recipe_body(tips=None) == ""
    assert render_structured_recipe_body(ingredients=[], steps=[], tips=[]) == ""


def test_structured_tips_one():
    html = render_structured_recipe_body(tips=["夏天水温要更低"])
    assert html == (
        '<ul class="recipe-tips">'
        '<li class="recipe-tips-item">夏天水温要更低</li>'
        '</ul>'
    )


def test_structured_tips_many():
    html = render_structured_recipe_body(tips=["夏天水温要更低", "饧面不要超过 20 分钟", "按口味调整盐"])
    assert html == (
        '<ul class="recipe-tips">'
        '<li class="recipe-tips-item">夏天水温要更低</li>'
        '<li class="recipe-tips-item">饧面不要超过 20 分钟</li>'
        '<li class="recipe-tips-item">按口味调整盐</li>'
        '</ul>'
    )


def test_structured_tips_escapes_html():
    html = render_structured_recipe_body(tips=['<script>x</script> 1<"'])
    assert "<script>" not in html
    assert "&lt;script&gt;" in html
    assert "&quot;" in html


def test_structured_body_concatenates_ingredients_then_steps_then_tips():
    html = render_structured_recipe_body(
        ingredients=[{"name": "面粉", "amount": "200", "unit": "g"}],
        steps=["混合面粉与水"],
        tips=["夏天水温要更低"],
    )
    assert html == (
        '<table class="recipe-ingredients"><tbody>'
        '<tr>'
        '<td class="recipe-ingredients-name">面粉</td>'
        '<td class="recipe-ingredients-amount">200 g</td>'
        '</tr>'
        '</tbody></table>'
        '<ol class="recipe-steps">'
        '<li class="recipe-steps-item">混合面粉与水</li>'
        '</ol>'
        '<ul class="recipe-tips">'
        '<li class="recipe-tips-item">夏天水温要更低</li>'
        '</ul>'
    )


def test_html_keeps_structured_tip_list_classes():
    html = render_markdown_to_html(
        '# T\n\n## S\n\n### 面团\n\n'
        '<ul class="recipe-tips">'
        '<li class="recipe-tips-item">夏天水温要更低</li>'
        '</ul>\n'
    )
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    ul = soup.select_one("ul.recipe-tips")
    assert ul is not None
    assert ul.select_one("li.recipe-tips-item").get_text() == "夏天水温要更低"


def test_station_html_tips_only_uses_structured_seam():
    html = render_station_html("肠粉档", [
        ParsedRecipe(
            section="配方", recipe_name="面团",
            body_markdown="这段 Markdown 不应出现在结构化卡片里",
            sort_order=0, id=11,
            tips=("夏天水温要更低", "饧面不要超过 20 分钟"),
        ),
    ])
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    card = soup.select_one("article.recipe-card")
    assert card is not None
    ul = card.select_one("ul.recipe-tips")
    assert ul is not None
    items = [li.get_text() for li in ul.select("li.recipe-tips-item")]
    assert items == ["夏天水温要更低", "饧面不要超过 20 分钟"]
    assert "这段 Markdown 不应出现在结构化卡片里" not in html
    assert card.select_one(".recipe-ingredients") is None
    assert card.select_one(".recipe-steps") is None


def test_station_html_keeps_ingredients_and_steps_when_tips_present():
    html = render_station_html("肠粉档", [
        ParsedRecipe(
            section="配方", recipe_name="面团", body_markdown="",
            sort_order=0, id=12,
            ingredients=({"name": "面粉", "amount": "200", "unit": "g"},),
            steps=("混合面粉与水",),
            tips=("夏天水温要更低",),
        ),
    ])
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    body = soup.select_one(".recipe-card-body")
    assert body.select_one("table.recipe-ingredients") is not None
    assert body.select_one("ol.recipe-steps") is not None
    assert body.select_one("ul.recipe-tips") is not None
    assert body.select_one(".recipe-ingredients-name").get_text() == "面粉"
    assert body.select_one(".recipe-steps-item").get_text() == "混合面粉与水"
    assert body.select_one(".recipe-tips-item").get_text() == "夏天水温要更低"


def test_station_html_stamps_base_servings_when_qty_present():
    html = render_station_html("肠粉档", [
        ParsedRecipe(
            section="配方", recipe_name="面团", body_markdown="",
            sort_order=0, id=13,
            ingredients=({"name": "面粉", "amount": "200", "unit": "g"},),
            base_servings_qty=4.5, base_servings_unit="人份",
        ),
    ])
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    card = soup.select_one("article.recipe-card")
    assert card.get("data-base-servings-qty") == "4.5"
    assert card.get("data-base-servings-unit") == "人份"
    h3 = card.select_one("h3")
    assert h3.get("data-base-servings-qty") == "4.5"
    assert h3.get("data-base-servings-unit") == "人份"


def test_station_html_omits_base_servings_when_qty_null():
    html = render_station_html("肠粉档", [
        ParsedRecipe(
            section="配方", recipe_name="面团", body_markdown="x",
            sort_order=0, id=14,
        ),
    ])
    assert "data-base-servings-qty" not in html
    assert "data-base-servings-unit" not in html


def test_station_html_omits_base_servings_when_qty_not_positive():
    html = render_station_html("肠粉档", [
        ParsedRecipe(
            section="配方", recipe_name="面团", body_markdown="x",
            sort_order=0, id=15,
            base_servings_qty=0, base_servings_unit="人份",
        ),
    ])
    assert "data-base-servings-qty" not in html
    assert "data-base-servings-unit" not in html

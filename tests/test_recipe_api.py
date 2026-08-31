import asyncio
import csv
import io
import sqlite3

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

import api.recipes as recipes_module
from services.recipes.store import RecipeStore
from config import settings
from database import DatabaseManager
from services import auth_service
from services.app_runtime import AppRuntime, set_runtime


def _get_loop():
    try:
        return asyncio.get_event_loop()
    except RuntimeError:
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        return loop


def _run(coro):
    return _get_loop().run_until_complete(coro)


def _seed(path):
    conn = sqlite3.connect(path)
    conn.executescript(
        """
        CREATE TABLE sop_stations (slug TEXT PRIMARY KEY, title TEXT NOT NULL, updated_at TEXT NOT NULL);
        CREATE TABLE sop_recipes (
            id INTEGER PRIMARY KEY AUTOINCREMENT, station_slug TEXT NOT NULL, section TEXT NOT NULL,
            recipe_name TEXT NOT NULL, body_markdown TEXT NOT NULL, sort_order INTEGER NOT NULL,
            is_new INTEGER NOT NULL DEFAULT 0, is_active INTEGER NOT NULL DEFAULT 1, updated_at TEXT NOT NULL);
        INSERT INTO sop_stations VALUES ('changfen','肠粉档','2026-01-01T00:00:00+00:00');
        INSERT INTO sop_recipes (station_slug,section,recipe_name,body_markdown,sort_order,is_new,is_active,updated_at)
        VALUES ('changfen','配方','肠粉酱油','酱油：100g',0,0,1,'2026-01-01T00:00:00+00:00');
        """
    )
    conn.commit()
    conn.close()


@pytest.fixture
def recipe_store(tmp_path):
    db_path = tmp_path / "recipes.db"
    _seed(str(db_path))
    store = RecipeStore(str(db_path))
    _run(store.connect())
    yield store
    _run(store.close())


@pytest.fixture
def client(recipe_store, tmp_path):
    store = recipe_store

    old_database_dir = settings.DATABASE_DIR
    settings.DATABASE_DIR = str(tmp_path)
    auth_db = DatabaseManager()
    _run(auth_db.connect())
    set_runtime(AppRuntime(db=auth_db))
    _run(auth_service.init_user("admin", "password123"))
    token, _meta = _run(auth_service.issue_api_token(label="test-recipe-api"))

    app = FastAPI()
    app.include_router(recipes_module.router)
    app.dependency_overrides[recipes_module._get_recipe_store] = lambda: store
    with TestClient(app) as c:
        c.headers.update({"X-Admin-Token": token})
        yield c

    _run(auth_db.close())
    set_runtime(None)
    settings.DATABASE_DIR = old_database_dir


def test_list_stations(client):
    r = client.get("/api/recipes/stations")
    assert r.status_code == 200
    station = r.json()["stations"][0]
    assert station["slug"] == "changfen"
    assert station["title"] == "肠粉档"


def test_station_detail_fragment(client):
    r = client.get("/api/recipes/stations/changfen")
    assert r.status_code == 200
    data = r.json()
    assert data["title"] == "肠粉档"
    assert "sop-doc" in data["content_html"]


def test_station_detail_stamps_recipe_id_on_cards(client):
    html = client.get("/api/recipes/stations/changfen").json()["content_html"]
    assert 'data-recipe-id="1"' in html
    assert "recipe-card" in html


def test_station_detail_stamps_base_servings_on_cards(client):
    rid = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团份数",
            "body": "",
            "is_new": False,
            "ingredients": [{"name": "面粉", "amount": "200", "unit": "g"}],
            "base_servings_qty": 4.5,
            "base_servings_unit": "人份",
        },
    ).json()["id"]
    html = client.get("/api/recipes/stations/changfen").json()["content_html"]
    assert f'data-recipe-id="{rid}"' in html
    assert 'data-base-servings-qty="4.5"' in html
    assert 'data-base-servings-unit="人份"' in html


def test_station_detail_omits_base_servings_when_null(client):
    html = client.get("/api/recipes/stations/changfen").json()["content_html"]
    # seed recipe has no base servings
    assert "data-base-servings-qty" not in html
    assert "data-base-servings-unit" not in html


def test_station_detail_renders_structured_ingredients_table(client):
    rid = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "这段 Markdown 不应出现在结构化卡片里",
            "is_new": False,
            "ingredients": [{"name": "面粉", "amount": "200", "unit": "g"}],
        },
    ).json()["id"]
    html = client.get("/api/recipes/stations/changfen").json()["content_html"]
    assert "recipe-ingredients" in html
    assert "面粉" in html
    assert "200 g" in html
    assert f'data-recipe-id="{rid}"' in html
    assert "这段 Markdown 不应出现在结构化卡片里" not in html
    assert "酱油：100g" in html
    assert "sop-section-grid" in html


def test_station_detail_print_html_includes_structured_card_classes(client):
    client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团全套",
            "body": "",
            "is_new": False,
            "ingredients": [{"name": "面粉", "amount": "200", "unit": "g"}],
            "steps": ["混合面粉与水"],
            "tips": ["夏天水温要更低"],
        },
    )
    html = client.get("/api/recipes/stations/changfen").json()["content_html"]
    assert "recipe-card" in html
    assert "recipe-ingredients" in html
    assert "recipe-steps" in html
    assert "recipe-tips" in html


def test_station_detail_renders_structured_steps_list(client):
    rid = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团做法",
            "body": "这段 Markdown 不应出现在步骤卡片里",
            "is_new": False,
            "steps": ["混合面粉与水", "静置 10 分钟"],
        },
    ).json()["id"]
    html = client.get("/api/recipes/stations/changfen").json()["content_html"]
    assert "recipe-steps" in html
    assert "recipe-steps-item" in html
    assert "混合面粉与水" in html
    assert "静置 10 分钟" in html
    assert f'data-recipe-id="{rid}"' in html
    assert "这段 Markdown 不应出现在步骤卡片里" not in html
    assert "酱油：100g" in html


def test_station_detail_404(client):
    assert client.get("/api/recipes/stations/nope").status_code == 404


def test_create_recipe(client):
    r = client.post("/api/recipes/stations/changfen/recipes",
                    json={"section": "配方", "recipe_name": "新", "body": "x", "is_new": False})
    assert r.status_code == 200
    assert isinstance(r.json()["id"], int)


def test_create_recipe_validation(client):
    r = client.post("/api/recipes/stations/changfen/recipes",
                    json={"section": "配方", "recipe_name": "", "body": "x"})
    assert r.status_code == 400
    assert r.json()["detail"] == "配方名称不能为空"


def test_create_recipe_unknown_section_rejected(client):
    r = client.post("/api/recipes/stations/changfen/recipes",
                    json={"section": "粥品", "recipe_name": "新", "body": "x"})
    assert r.status_code == 400
    assert r.json()["detail"] == "章节必须是配方、出品标准、检核要求或食安要求"
    r = client.post("/api/recipes/stations/changfen/recipes",
                    json={"section": "  ", "recipe_name": "新", "body": "x"})
    assert r.status_code == 400
    assert r.json()["detail"] == "章节不能为空"


def test_missing_recipe_uses_recipe_word(client):
    r = client.get("/api/recipes/recipes/99999/history")
    assert r.status_code == 404
    assert r.json()["detail"] == "配方不存在"


def test_export_csv(client):
    r = client.get("/api/recipes/stations/changfen/export")
    assert r.status_code == 200
    assert "text/csv" in r.headers["content-type"]
    assert "section,recipe_name" in r.text


def test_export_csv_includes_structured_json_columns(client):
    client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "兜底正文",
            "is_new": False,
            "ingredients": [{"name": "面粉", "amount": "200", "unit": "g"}],
            "steps": ["混合面粉与水"],
            "tips": ["夏天水温要更低"],
        },
    )
    r = client.get("/api/recipes/stations/changfen/export")
    assert r.status_code == 200
    reader = csv.DictReader(io.StringIO(r.text.lstrip("\ufeff")))
    assert reader.fieldnames == [
        "section", "recipe_name", "body_markdown", "sort_order", "is_new",
        "ingredients_json", "steps_json", "tips_json",
    ]
    dough = next(row for row in reader if row["recipe_name"] == "面团")
    assert dough["body_markdown"] == "兜底正文"
    assert dough["ingredients_json"] == '[{"name": "面粉", "amount": "200", "unit": "g"}]'
    assert dough["steps_json"] == '["混合面粉与水"]'
    assert dough["tips_json"] == '["夏天水温要更低"]'


def test_csv_export_import_roundtrip_structured_fields(client):
    client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "兜底正文",
            "is_new": False,
            "ingredients": [{"name": "面粉", "amount": "200", "unit": "g"}],
            "steps": ["混合面粉与水"],
            "tips": ["夏天水温要更低"],
        },
    )
    exported = client.get("/api/recipes/stations/changfen/export")
    assert exported.status_code == 200
    created = client.post("/api/recipes/stations", json={"slug": "xibing", "title": "西饼档"})
    assert created.status_code == 200
    imported = client.post(
        "/api/recipes/stations/xibing/import",
        files={"csv_file": ("recipes.csv", exported.content, "text/csv")},
    )
    assert imported.status_code == 200
    recipes = client.get("/api/recipes/stations/xibing/recipes").json()["recipes"]
    dough = next(row for row in recipes if row["recipe_name"] == "面团")
    assert dough["ingredients"] == [{"name": "面粉", "amount": "200", "unit": "g"}]
    assert dough["steps"] == ["混合面粉与水"]
    assert dough["tips"] == ["夏天水温要更低"]
    assert dough["body_markdown"] == "兜底正文"


def test_import_csv_invalid_ingredients_json_returns_400(client):
    csv_body = (
        "section,recipe_name,body_markdown,sort_order,is_new,"
        "ingredients_json,steps_json,tips_json\n"
        "配方,面团,x,0,0,not-json,[],[]\n"
    )
    r = client.post(
        "/api/recipes/stations/changfen/import",
        files={"csv_file": ("recipes.csv", csv_body.encode("utf-8"), "text/csv")},
    )
    assert r.status_code == 400
    assert "第 2 行" in r.json()["detail"]
    assert "用料" in r.json()["detail"]


def test_docx_export(client):
    r = client.get("/api/recipes/stations/changfen/docx")
    assert r.status_code == 200
    assert "officedocument" in r.headers["content-type"]


def test_docx_export_includes_structured_fields(client):
    from docx import Document

    client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "这段 Markdown 不应出现",
            "is_new": False,
            "ingredients": [{"name": "面粉", "amount": "200", "unit": "g"}],
            "steps": ["混合面粉与水"],
            "tips": ["夏天水温要更低"],
        },
    )
    r = client.get("/api/recipes/stations/changfen/docx")
    assert r.status_code == 200
    doc = Document(io.BytesIO(r.content))
    cell_texts = [cell.text.strip() for table in doc.tables for row in table.rows for cell in row.cells]
    assert "面粉" in cell_texts
    assert any("200" in text for text in cell_texts)
    paragraph_texts = [p.text for p in doc.paragraphs]
    assert any("混合面粉与水" in text for text in paragraph_texts)
    assert any("夏天水温要更低" in text for text in paragraph_texts)
    assert not any("这段 Markdown 不应出现" in text for text in paragraph_texts)


def test_station_detail_include_inactive(client):
    r = client.post("/api/recipes/stations/changfen/recipes",
                    json={"section": "配方", "recipe_name": "停售品", "body": "y"})
    rid_off = r.json()["id"]
    client.post(f"/api/recipes/recipes/{rid_off}/toggle-active")  # 置停用
    d = client.get("/api/recipes/stations/changfen").json()
    assert "停售品" not in d["content_html"]
    d2 = client.get("/api/recipes/stations/changfen?include_inactive=1").json()
    assert "停售品" in d2["content_html"]
    assert "recipe-card--inactive" in d2["content_html"]


def test_import_csv_does_not_infer_is_new_from_name_or_body(client):
    csv_body = (
        "section,recipe_name,body_markdown,sort_order,is_new\n"
        "配方,【新】菠菜饺,正文含【新】标记,10,0\n"
    )
    r = client.post(
        "/api/recipes/stations/changfen/import",
        files={"csv_file": ("recipes.csv", csv_body.encode("utf-8"), "text/csv")},
    )
    assert r.status_code == 200
    recipes = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"]
    imported = next(row for row in recipes if row["recipe_name"] == "【新】菠菜饺")
    assert imported["is_new"] == 0


def test_reorder_recipes_requires_auth(client):
    client.headers.pop("X-Admin-Token", None)
    r = client.put("/api/recipes/stations/changfen/recipes/reorder", json={"ids": [1]})
    assert r.status_code == 401


def test_reorder_recipes_assigns_unique_sort_order(client):
    first_id = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"][0]["id"]
    second_id = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={"section": "配方", "recipe_name": "第二", "body": "b", "is_new": False},
    ).json()["id"]
    third_id = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={"section": "出品标准", "recipe_name": "第三", "body": "c", "is_new": False},
    ).json()["id"]

    r = client.put(
        "/api/recipes/stations/changfen/recipes/reorder",
        json={"ids": [third_id, first_id, second_id]},
    )
    assert r.status_code == 200

    recipes = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"]
    assert [row["id"] for row in recipes] == [third_id, first_id, second_id]
    assert [row["sort_order"] for row in recipes] == [0, 1, 2]
    history = client.get(f"/api/recipes/recipes/{first_id}/history").json()["history"]
    assert history == []


def test_search_empty_query_returns_empty_groups(client):
    r = client.get("/api/recipes/search")
    assert r.status_code == 200
    assert r.json() == {"groups": []}

    r2 = client.get("/api/recipes/search", params={"q": ""})
    assert r2.status_code == 200
    assert r2.json() == {"groups": []}

    r3 = client.get("/api/recipes/search", params={"q": "   "})
    assert r3.status_code == 200
    assert r3.json() == {"groups": []}


def test_search_groups_matches_by_station(client):
    client.post("/api/recipes/stations", json={"slug": "shulong", "title": "熟笼档"})
    dumpling = client.post(
        "/api/recipes/stations/shulong/recipes",
        json={"section": "配方", "recipe_name": "鲜虾饺", "body": "虾馅", "is_new": False},
    ).json()["id"]
    sauce = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"][0]
    assert sauce["recipe_name"] == "肠粉酱油"

    r = client.get("/api/recipes/search", params={"q": "虾"})
    assert r.status_code == 200
    assert r.json() == {
        "groups": [
            {
                "station_slug": "shulong",
                "station_title": "熟笼档",
                "items": [
                    {"recipe_id": dumpling, "recipe_name": "鲜虾饺", "section": "配方"},
                ],
            },
        ],
    }

    r2 = client.get("/api/recipes/search", params={"q": "肠粉"})
    assert r2.json() == {
        "groups": [
            {
                "station_slug": "changfen",
                "station_title": "肠粉档",
                "items": [
                    {"recipe_id": sauce["id"], "recipe_name": "肠粉酱油", "section": "配方"},
                ],
            },
        ],
    }


def test_search_returns_multiple_items_in_one_station(client):
    sauce = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"][0]
    check_id = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={"section": "出品标准", "recipe_name": "肠粉检核", "body": "x", "is_new": False},
    ).json()["id"]
    items = client.get("/api/recipes/search", params={"q": "肠粉"}).json()["groups"][0]["items"]
    assert items == [
        {"recipe_id": sauce["id"], "recipe_name": "肠粉酱油", "section": "配方"},
        {"recipe_id": check_id, "recipe_name": "肠粉检核", "section": "出品标准"},
    ]


def test_search_is_case_insensitive_and_skips_body(client):
    wrap_id = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={"section": "配方", "recipe_name": "Spinach Wrap", "body": "hidden dumpling filling", "is_new": False},
    ).json()["id"]
    decoy_id = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={"section": "配方", "recipe_name": "无关名字", "body": "Spinach Wrap 写在正文", "is_new": False},
    ).json()["id"]

    r = client.get("/api/recipes/search", params={"q": "spinach wrap"})
    names = [item["recipe_name"] for g in r.json()["groups"] for item in g["items"]]
    ids = [item["recipe_id"] for g in r.json()["groups"] for item in g["items"]]
    assert names == ["Spinach Wrap"]
    assert ids == [wrap_id]
    assert decoy_id not in ids

    r2 = client.get("/api/recipes/search", params={"q": "hidden dumpling"})
    assert r2.json() == {"groups": []}


def test_search_excludes_inactive_unless_flagged(client):
    r = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={"section": "配方", "recipe_name": "停用虾饺", "body": "y", "is_new": False},
    )
    rid = r.json()["id"]
    client.post(f"/api/recipes/recipes/{rid}/toggle-active")

    assert client.get("/api/recipes/search", params={"q": "停用虾饺"}).json() == {"groups": []}

    flagged = client.get("/api/recipes/search", params={"q": "停用虾饺", "include_inactive": 1}).json()
    assert flagged["groups"][0]["items"] == [
        {"recipe_id": rid, "recipe_name": "停用虾饺", "section": "配方"},
    ]


def test_search_treats_like_wildcards_as_literal(client):
    percent_id = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={"section": "配方", "recipe_name": "50%糖浆", "body": "x", "is_new": False},
    ).json()["id"]
    client.post(
        "/api/recipes/stations/changfen/recipes",
        json={"section": "配方", "recipe_name": "普通糖浆", "body": "x", "is_new": False},
    )

    wildcard = client.get("/api/recipes/search", params={"q": "%"}).json()
    assert wildcard == {
        "groups": [
            {
                "station_slug": "changfen",
                "station_title": "肠粉档",
                "items": [
                    {"recipe_id": percent_id, "recipe_name": "50%糖浆", "section": "配方"},
                ],
            },
        ],
    }

    underscore = client.get("/api/recipes/search", params={"q": "_"}).json()
    assert underscore == {"groups": []}


def test_search_stays_public(client):
    client.headers.pop("X-Admin-Token", None)
    r = client.get("/api/recipes/search", params={"q": "肠粉"})
    assert r.status_code == 200
    assert r.json()["groups"][0]["items"][0]["recipe_name"] == "肠粉酱油"


def test_create_omitting_ingredients_still_works(client):
    r = client.post("/api/recipes/stations/changfen/recipes",
                    json={"section": "配方", "recipe_name": "无用料", "body": "x", "is_new": False})
    assert r.status_code == 200
    rid = r.json()["id"]
    recipes = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"]
    row = next(x for x in recipes if x["id"] == rid)
    assert row["ingredients"] == []
    current = client.get(f"/api/recipes/recipes/{rid}/history").json()["current"]
    assert current["ingredients"] == []


def test_create_and_list_recipes_include_ingredients(client):
    r = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "",
            "is_new": False,
            "ingredients": [
                {"name": "面粉", "amount": "200", "unit": "g"},
                {"name": "水", "amount": "适量", "unit": ""},
            ],
        },
    )
    assert r.status_code == 200
    rid = r.json()["id"]
    recipes = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"]
    row = next(x for x in recipes if x["id"] == rid)
    assert row["ingredients"] == [
        {"name": "面粉", "amount": "200", "unit": "g"},
        {"name": "水", "amount": "适量", "unit": ""},
    ]
    current = client.get(f"/api/recipes/recipes/{rid}/history").json()["current"]
    assert current["ingredients"] == row["ingredients"]


def test_update_recipe_ingredients(client):
    rid = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"][0]["id"]
    r = client.put(
        f"/api/recipes/recipes/{rid}",
        json={
            "section": "配方",
            "recipe_name": "肠粉酱油",
            "body": "酱油：100g",
            "is_new": False,
            "ingredients": [{"name": "酱油", "amount": "100", "unit": "g"}],
        },
    )
    assert r.status_code == 200
    assert r.json()["ingredients"] == [{"name": "酱油", "amount": "100", "unit": "g"}]


def test_update_omitting_ingredients_still_works(client):
    rid = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"][0]["id"]
    r = client.put(
        f"/api/recipes/recipes/{rid}",
        json={"section": "配方", "recipe_name": "改名", "body": "新正文", "is_new": False},
    )
    assert r.status_code == 200
    assert r.json()["recipe_name"] == "改名"
    assert r.json()["ingredients"] == []


def test_update_omitting_ingredients_preserves_existing_rows(client):
    rid = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "",
            "ingredients": [{"name": "面粉", "amount": "200", "unit": "g"}],
        },
    ).json()["id"]
    r = client.put(
        f"/api/recipes/recipes/{rid}",
        json={"section": "配方", "recipe_name": "面团改", "body": "", "is_new": False},
    )
    assert r.status_code == 200
    assert r.json()["recipe_name"] == "面团改"
    assert r.json()["ingredients"] == [{"name": "面粉", "amount": "200", "unit": "g"}]


def test_ingredient_name_over_limit_returns_400(client):
    r = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "",
            "ingredients": [{"name": "a" * 121, "amount": "1", "unit": "g"}],
        },
    )
    assert r.status_code == 400
    assert r.json()["detail"] == "用料名称长度不能超过 120 个字符"


def test_ingredient_amount_over_limit_returns_400(client):
    rid = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"][0]["id"]
    r = client.put(
        f"/api/recipes/recipes/{rid}",
        json={
            "section": "配方",
            "recipe_name": "肠粉酱油",
            "body": "",
            "ingredients": [{"name": "酱油", "amount": "b" * 121, "unit": "g"}],
        },
    )
    assert r.status_code == 400
    assert r.json()["detail"] == "用料用量长度不能超过 120 个字符"


def test_ingredient_unit_over_limit_returns_400(client):
    r = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "",
            "ingredients": [{"name": "盐", "amount": "1", "unit": "c" * 121}],
        },
    )
    assert r.status_code == 400
    assert r.json()["detail"] == "用料单位长度不能超过 120 个字符"


def test_create_omitting_steps_still_works(client):
    r = client.post("/api/recipes/stations/changfen/recipes",
                    json={"section": "配方", "recipe_name": "无步骤", "body": "x", "is_new": False})
    assert r.status_code == 200
    rid = r.json()["id"]
    recipes = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"]
    row = next(x for x in recipes if x["id"] == rid)
    assert row["steps"] == []
    current = client.get(f"/api/recipes/recipes/{rid}/history").json()["current"]
    assert current["steps"] == []


def test_create_and_list_recipes_include_steps(client):
    r = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "",
            "is_new": False,
            "steps": ["混合面粉与水", "静置 10 分钟"],
        },
    )
    assert r.status_code == 200
    rid = r.json()["id"]
    recipes = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"]
    row = next(x for x in recipes if x["id"] == rid)
    assert row["steps"] == ["混合面粉与水", "静置 10 分钟"]
    current = client.get(f"/api/recipes/recipes/{rid}/history").json()["current"]
    assert current["steps"] == row["steps"]


def test_update_recipe_steps(client):
    rid = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"][0]["id"]
    r = client.put(
        f"/api/recipes/recipes/{rid}",
        json={
            "section": "配方",
            "recipe_name": "肠粉酱油",
            "body": "酱油：100g",
            "is_new": False,
            "steps": ["把酱油煮开"],
        },
    )
    assert r.status_code == 200
    assert r.json()["steps"] == ["把酱油煮开"]


def test_update_omitting_steps_preserves_existing_rows(client):
    rid = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "",
            "steps": ["混合面粉与水"],
        },
    ).json()["id"]
    r = client.put(
        f"/api/recipes/recipes/{rid}",
        json={"section": "配方", "recipe_name": "面团改", "body": "", "is_new": False},
    )
    assert r.status_code == 200
    assert r.json()["recipe_name"] == "面团改"
    assert r.json()["steps"] == ["混合面粉与水"]


def test_step_text_over_limit_returns_400(client):
    r = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "",
            "steps": ["a" * 121],
        },
    )
    assert r.status_code == 400
    assert r.json()["detail"] == "步骤长度不能超过 120 个字符"


def test_steps_count_over_limit_returns_400(client):
    r = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "",
            "steps": ["步"] * 51,
        },
    )
    assert r.status_code == 400
    assert r.json()["detail"] == "步骤数量不能超过 50 个"


def test_station_detail_renders_structured_tips_list(client):
    rid = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团贴士",
            "body": "这段 Markdown 不应出现在贴士卡片里",
            "is_new": False,
            "tips": ["夏天水温要更低", "饧面不要超过 20 分钟"],
        },
    ).json()["id"]
    html = client.get("/api/recipes/stations/changfen").json()["content_html"]
    assert "recipe-tips" in html
    assert "recipe-tips-item" in html
    assert "夏天水温要更低" in html
    assert "饧面不要超过 20 分钟" in html
    assert f'data-recipe-id="{rid}"' in html
    assert "这段 Markdown 不应出现在贴士卡片里" not in html
    assert "酱油：100g" in html


def test_create_omitting_tips_still_works(client):
    r = client.post("/api/recipes/stations/changfen/recipes",
                    json={"section": "配方", "recipe_name": "无贴士", "body": "x", "is_new": False})
    assert r.status_code == 200
    rid = r.json()["id"]
    recipes = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"]
    row = next(x for x in recipes if x["id"] == rid)
    assert row["tips"] == []
    current = client.get(f"/api/recipes/recipes/{rid}/history").json()["current"]
    assert current["tips"] == []


def test_create_and_list_recipes_include_tips(client):
    r = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "",
            "is_new": False,
            "tips": ["夏天水温要更低", "饧面不要超过 20 分钟"],
        },
    )
    assert r.status_code == 200
    rid = r.json()["id"]
    recipes = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"]
    row = next(x for x in recipes if x["id"] == rid)
    assert row["tips"] == ["夏天水温要更低", "饧面不要超过 20 分钟"]
    current = client.get(f"/api/recipes/recipes/{rid}/history").json()["current"]
    assert current["tips"] == row["tips"]


def test_update_recipe_tips(client):
    rid = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"][0]["id"]
    r = client.put(
        f"/api/recipes/recipes/{rid}",
        json={
            "section": "配方",
            "recipe_name": "肠粉酱油",
            "body": "酱油：100g",
            "is_new": False,
            "tips": ["按口味调整盐"],
        },
    )
    assert r.status_code == 200
    assert r.json()["tips"] == ["按口味调整盐"]


def test_update_omitting_tips_preserves_existing_rows(client):
    rid = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "",
            "tips": ["夏天水温要更低"],
        },
    ).json()["id"]
    r = client.put(
        f"/api/recipes/recipes/{rid}",
        json={"section": "配方", "recipe_name": "面团改", "body": "", "is_new": False},
    )
    assert r.status_code == 200
    assert r.json()["recipe_name"] == "面团改"
    assert r.json()["tips"] == ["夏天水温要更低"]


def test_tip_text_over_limit_returns_400(client):
    r = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "",
            "tips": ["a" * 121],
        },
    )
    assert r.status_code == 400
    assert r.json()["detail"] == "小贴士长度不能超过 120 个字符"


def test_tips_count_over_limit_returns_400(client):
    r = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "",
            "tips": ["贴"] * 51,
        },
    )
    assert r.status_code == 400
    assert r.json()["detail"] == "小贴士数量不能超过 50 个"


def test_create_omitting_base_servings_defaults_null(client):
    r = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={"section": "配方", "recipe_name": "无份数", "body": "x", "is_new": False},
    )
    assert r.status_code == 200
    rid = r.json()["id"]
    recipes = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"]
    row = next(x for x in recipes if x["id"] == rid)
    assert row["base_servings_qty"] is None
    assert row["base_servings_unit"] is None
    current = client.get(f"/api/recipes/recipes/{rid}/history").json()["current"]
    assert current["base_servings_qty"] is None
    assert current["base_servings_unit"] is None


def test_create_and_list_recipes_include_base_servings(client):
    r = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "",
            "is_new": False,
            "base_servings_qty": 4.5,
            "base_servings_unit": "人份",
        },
    )
    assert r.status_code == 200
    rid = r.json()["id"]
    recipes = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"]
    row = next(x for x in recipes if x["id"] == rid)
    assert row["base_servings_qty"] == 4.5
    assert row["base_servings_unit"] == "人份"
    current = client.get(f"/api/recipes/recipes/{rid}/history").json()["current"]
    assert current["base_servings_qty"] == 4.5
    assert current["base_servings_unit"] == "人份"


def test_update_recipe_base_servings(client):
    rid = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"][0]["id"]
    r = client.put(
        f"/api/recipes/recipes/{rid}",
        json={
            "section": "配方",
            "recipe_name": "肠粉酱油",
            "body": "酱油：100g",
            "is_new": False,
            "base_servings_qty": 2,
            "base_servings_unit": "份",
        },
    )
    assert r.status_code == 200
    assert r.json()["base_servings_qty"] == 2
    assert r.json()["base_servings_unit"] == "份"


def test_update_omitting_base_servings_preserves_existing(client):
    rid = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "",
            "base_servings_qty": 4,
            "base_servings_unit": "人份",
        },
    ).json()["id"]
    r = client.put(
        f"/api/recipes/recipes/{rid}",
        json={"section": "配方", "recipe_name": "面团改", "body": "", "is_new": False},
    )
    assert r.status_code == 200
    assert r.json()["recipe_name"] == "面团改"
    assert r.json()["base_servings_qty"] == 4
    assert r.json()["base_servings_unit"] == "人份"


def test_update_zero_base_servings_qty_stores_null(client):
    rid = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "",
            "base_servings_qty": 4,
            "base_servings_unit": "人份",
        },
    ).json()["id"]
    r = client.put(
        f"/api/recipes/recipes/{rid}",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "",
            "is_new": False,
            "base_servings_qty": 0,
            "base_servings_unit": "人份",
        },
    )
    assert r.status_code == 200
    assert r.json()["base_servings_qty"] is None
    assert r.json()["base_servings_unit"] is None


def test_negative_base_servings_qty_returns_400(client):
    r = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "",
            "base_servings_qty": -1,
            "base_servings_unit": "人份",
        },
    )
    assert r.status_code == 400
    assert r.json()["detail"] == "基准份数不能为负数"


def test_update_negative_base_servings_qty_returns_400(client):
    rid = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"][0]["id"]
    r = client.put(
        f"/api/recipes/recipes/{rid}",
        json={
            "section": "配方",
            "recipe_name": "肠粉酱油",
            "body": "",
            "is_new": False,
            "base_servings_qty": -0.5,
        },
    )
    assert r.status_code == 400
    assert r.json()["detail"] == "基准份数不能为负数"


def _flag_review(store, recipe_id, *, needs_review=1, legacy_markdown="旧正文\n第二行"):
    conn = sqlite3.connect(store.db_path)
    conn.execute(
        "UPDATE sop_recipes SET needs_review=?, legacy_markdown=? WHERE id=?",
        (needs_review, legacy_markdown, recipe_id),
    )
    conn.commit()
    conn.close()


def test_list_get_history_include_needs_review_and_legacy_markdown(client, recipe_store):
    rid = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"][0]["id"]
    listed = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"][0]
    assert listed["needs_review"] == 0
    assert listed["legacy_markdown"] is None
    current = client.get(f"/api/recipes/recipes/{rid}/history").json()["current"]
    assert current["needs_review"] == 0
    assert current["legacy_markdown"] is None

    snapshot = "面粉 200g\n混合"
    _flag_review(recipe_store, rid, legacy_markdown=snapshot)
    listed = next(
        row for row in client.get("/api/recipes/stations/changfen/recipes").json()["recipes"]
        if row["id"] == rid
    )
    assert listed["needs_review"] == 1
    assert listed["legacy_markdown"] == snapshot
    current = client.get(f"/api/recipes/recipes/{rid}/history").json()["current"]
    assert current["needs_review"] == 1
    assert current["legacy_markdown"] == snapshot


def test_list_stations_includes_needs_review_count(client, recipe_store):
    stations = client.get("/api/recipes/stations").json()["stations"]
    assert stations[0]["slug"] == "changfen"
    assert stations[0]["needs_review_count"] == 0

    rid = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"][0]["id"]
    _flag_review(recipe_store, rid)
    stations = client.get("/api/recipes/stations").json()["stations"]
    assert stations[0]["needs_review_count"] == 1


def test_confirm_review_sets_flag_zero_and_preserves_other_fields(client, recipe_store):
    created = client.post(
        "/api/recipes/stations/changfen/recipes",
        json={
            "section": "配方",
            "recipe_name": "面团",
            "body": "正文保留",
            "is_new": False,
            "ingredients": [{"name": "面粉", "amount": "200", "unit": "g"}],
            "steps": ["混合面粉与水"],
            "tips": ["夏天水温要更低"],
            "base_servings_qty": 4,
            "base_servings_unit": "人份",
        },
    )
    assert created.status_code == 200
    rid = created.json()["id"]
    snapshot = "迁移前原文\n盐 2g"
    _flag_review(recipe_store, rid, legacy_markdown=snapshot)

    r = client.post(f"/api/recipes/recipes/{rid}/confirm-review")
    assert r.status_code == 200
    body = r.json()
    assert body["needs_review"] == 0
    assert body["legacy_markdown"] == snapshot
    assert body["recipe_name"] == "面团"
    assert body["body_markdown"] == "正文保留"
    assert body["ingredients"] == [{"name": "面粉", "amount": "200", "unit": "g"}]
    assert body["steps"] == ["混合面粉与水"]
    assert body["tips"] == ["夏天水温要更低"]
    assert body["base_servings_qty"] == 4
    assert body["base_servings_unit"] == "人份"

    current = client.get(f"/api/recipes/recipes/{rid}/history").json()["current"]
    assert current["needs_review"] == 0
    assert current["legacy_markdown"] == snapshot
    assert current["ingredients"] == body["ingredients"]
    listed = next(
        row for row in client.get("/api/recipes/stations/changfen/recipes").json()["recipes"]
        if row["id"] == rid
    )
    assert listed["needs_review"] == 0
    assert listed["legacy_markdown"] == snapshot


def test_confirm_review_missing_returns_404(client):
    r = client.post("/api/recipes/recipes/99999/confirm-review")
    assert r.status_code == 404


def test_confirm_review_requires_auth(client):
    client.headers.pop("X-Admin-Token", None)
    r = client.post("/api/recipes/recipes/1/confirm-review")
    assert r.status_code == 401


def test_restore_history_applies_previous_version(client):
    rid = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"][0]["id"]
    client.put(
        f"/api/recipes/recipes/{rid}",
        json={
            "section": "配方",
            "recipe_name": "改名",
            "body": "新正文",
            "is_new": False,
            "ingredients": [{"name": "酱油", "amount": "80", "unit": "g"}],
            "steps": ["先改"],
            "tips": [],
        },
    )
    hist = client.get(f"/api/recipes/recipes/{rid}/history").json()["history"]
    assert hist[0]["recipe_name"] == "肠粉酱油"
    history_id = hist[0]["id"]
    restored = client.post(
        f"/api/recipes/recipes/{rid}/history/{history_id}/restore",
    )
    assert restored.status_code == 200
    assert restored.json()["recipe_name"] == "肠粉酱油"
    after = client.get(f"/api/recipes/recipes/{rid}/history").json()["history"]
    assert after[0]["recipe_name"] == "改名"
    assert after[0]["ingredients"] == [{"name": "酱油", "amount": "80", "unit": "g"}]


def test_restore_history_requires_auth(client):
    client.headers.pop("X-Admin-Token", None)
    r = client.post("/api/recipes/recipes/1/history/1/restore")
    assert r.status_code == 401


def test_restore_history_unknown_returns_404(client):
    rid = client.get("/api/recipes/stations/changfen/recipes").json()["recipes"][0]["id"]
    r = client.post(f"/api/recipes/recipes/{rid}/history/99999/restore")
    assert r.status_code == 404

